"""Load open_labels.xlsx into clean records.

- Normalises Cyrillic 'Х' and Latin 'X' ids to 'X-<n>'.
- Parses Russian timecodes like '22 секунда', '1минута 8 секунда'.
- Merges byte-identical videos (same md5) into one clip with all their timecodes.
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from engine.config import DATA_DIR, LABELS_XLSX


def norm_id(name: str) -> str:
    s = name.strip().replace("Х", "X").replace("х", "X").replace("x", "X")
    s = Path(s).stem
    m = re.match(r"X-?(\d+)$", s)
    if not m:
        raise ValueError(f"unexpected clip id: {name!r}")
    return f"X-{int(m.group(1))}"


def parse_timecode(text: str | None) -> float | None:
    if text is None or not str(text).strip():
        return None
    t = str(text).lower()
    mins = re.search(r"(\d+)\s*мин", t)
    secs = re.search(r"(\d+)\s*сек", t)
    if not mins and not secs:
        raise ValueError(f"cannot parse timecode: {text!r}")
    return float((int(mins.group(1)) * 60 if mins else 0) + (int(secs.group(1)) if secs else 0))


@dataclass
class Clip:
    clip_id: str                      # canonical id (first of any duplicates)
    path: Path
    md5: str
    violation: bool
    timecodes: list[float] = field(default_factory=list)
    aliases: list[str] = field(default_factory=list)  # all label ids that map to this file


def md5(path: Path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def video_index(data_dir: Path = DATA_DIR) -> dict[str, Path]:
    return {norm_id(p.name): p for p in sorted(data_dir.glob("*.mov"))}


def read_rows(xlsx: Path = LABELS_XLSX) -> list[tuple[str, float | None, bool]]:
    ws = load_workbook(xlsx, read_only=True).active
    rows = []
    for i, (name, tc, verdict) in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not name:
            continue
        is_viol = str(verdict).strip().lower().startswith("нарушение")
        rows.append((norm_id(str(name)), parse_timecode(tc), is_viol))
    return rows


def load_clips(xlsx: Path = LABELS_XLSX, data_dir: Path = DATA_DIR) -> list[Clip]:
    vids = video_index(data_dir)
    by_hash: dict[str, Clip] = {}
    for cid, tc, is_viol in read_rows(xlsx):
        path = vids[cid]
        h = md5(path)
        clip = by_hash.get(h)
        if clip is None:
            clip = by_hash[h] = Clip(cid, path, h, False)
        clip.aliases.append(cid)
        clip.violation |= is_viol
        if is_viol and tc is not None:
            clip.timecodes.append(tc)
    for c in by_hash.values():
        c.timecodes.sort()
    return sorted(by_hash.values(), key=lambda c: int(c.clip_id.split("-")[1]))


if __name__ == "__main__":
    clips = load_clips()
    for c in clips:
        print(c.clip_id, "aliases=" + ",".join(c.aliases), "V" if c.violation else "-", c.timecodes)
    print(len(clips), "unique clips;", sum(c.violation for c in clips), "with violations;",
          sum(len(c.timecodes) for c in clips), "timecodes")
